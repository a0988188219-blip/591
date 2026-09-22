"""Read listings to post from a plain Excel file.

This is the safe, no-scraping data source: you fill in one row per
property using data/listings_template.xlsx, and this module turns each
row into a Listing.
"""
from pathlib import Path

import openpyxl

from .models import Listing

COLUMNS = [
    "title", "city", "district", "address", "house_type",
    "total_price_wan", "main_building_ping", "total_ping", "land_ping",
    "rooms", "living_rooms", "bathrooms", "floor", "total_floors",
    "age_years", "parking", "facing", "description",
    "contact_name", "contact_phone", "photos",
]

EXAMPLE_ROW = [
    "高雄前鎮景觀三房", "高雄市", "前鎮區", "前鎮區中山二路OO號", "電梯大樓",
    880, 32.5, 38.2, None,
    3, 2, 2, 8, 15,
    12, "有(平面式)", "座北朝南", "採光佳、生活機能完善、鄰近捷運站，適合首購與換屋族。",
    "顏秀珊", "0913-335-182", "photos/sample1.jpg;photos/sample2.jpg",
]


def create_template(path: str | Path) -> None:
    """Generate an empty Excel template with headers + one example row."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "listings"
    ws.append(COLUMNS)
    ws.append(EXAMPLE_ROW)
    wb.save(path)


def load_listings(path: str | Path) -> list[Listing]:
    """Read every data row (skipping the header) into Listing objects."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    header = [c.value for c in ws[1]]
    col_index = {name: header.index(name) for name in COLUMNS if name in header}

    listings = []
    for row in rows:
        if row[0] is None:  # skip blank rows
            continue

        def get(col, default=None):
            idx = col_index.get(col)
            val = row[idx] if idx is not None else None
            return default if val is None else val

        photos_raw = get("photos", "")
        photos = [p.strip() for p in str(photos_raw).split(";") if p.strip()]

        listings.append(
            Listing(
                title=str(get("title", "")),
                city=str(get("city", "")),
                district=str(get("district", "")),
                address=str(get("address", "")),
                house_type=str(get("house_type", "")),
                total_price_wan=float(get("total_price_wan", 0) or 0),
                main_building_ping=float(get("main_building_ping", 0) or 0),
                total_ping=float(get("total_ping", 0) or 0),
                land_ping=float(get("land_ping")) if get("land_ping") not in (None, "") else None,
                rooms=int(get("rooms", 0) or 0),
                living_rooms=int(get("living_rooms", 0) or 0),
                bathrooms=int(get("bathrooms", 0) or 0),
                floor=int(get("floor", 0) or 0),
                total_floors=int(get("total_floors", 0) or 0),
                age_years=float(get("age_years", 0) or 0),
                parking=str(get("parking", "")),
                facing=str(get("facing", "")),
                description=str(get("description", "")),
                contact_name=str(get("contact_name", "")),
                contact_phone=str(get("contact_phone", "")),
                photos=photos,
            )
        )
    return listings
